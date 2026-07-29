#!/usr/bin/env python3
"""
extract-pptx.py — pull text, bullet levels, speaker notes, and images out of a
PowerPoint file (.pptx) into a clean JSON intermediate.

The JSON is NOT a finished deck. It is the raw material an agent maps into a
lewislulu-format HTML deck (.deck > .slide blocks, notes in <div class="notes">,
token-based CSS, runtime.js for nav). See:
    references/conversion-patterns.md  →  section "PPTX → deck"

Usage:
    python extract-pptx.py <input.pptx> [output_dir]

    output_dir defaults to ./<pptx-stem>-extract/ next to the .pptx.
    Images land in <output_dir>/images/ ; the JSON is <output_dir>/slides.json.

Requires: python-pptx
    pip install python-pptx

A venv with python-pptx already installed lives at:
    C:\\Users\\dougl\\Documents\\Claude Folder\\.venv-slides
On Windows:  C:\\Users\\dougl\\Documents\\Claude Folder\\.venv-slides\\Scripts\\python.exe extract-pptx.py deck.pptx
Any python with python-pptx on its path will also work — the venv is not required.
"""

import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.stderr.write(
        "error: python-pptx is not installed.\n"
        "  pip install python-pptx\n"
        "  (or use the venv at "
        "C:\\Users\\dougl\\Documents\\Claude Folder\\.venv-slides)\n"
    )
    sys.exit(2)


# EMU (English Metric Units) → px at 96 DPI. python-pptx returns sizes in EMU.
EMU_PER_INCH = 914400
PX_PER_INCH = 96


def emu_to_px(emu):
    if emu is None:
        return None
    return round(Emu(emu).inches * PX_PER_INCH)


def shape_top(shape):
    """Sort key so reading order roughly matches top-to-bottom, left-to-right."""
    try:
        return (shape.top if shape.top is not None else 0,
                shape.left if shape.left is not None else 0)
    except Exception:
        return (0, 0)


def extract_paragraphs(text_frame):
    """Return a list of {text, level} for each non-empty paragraph.

    `level` is the PowerPoint indent level (0 = top-level bullet). The agent
    uses it to nest <ul>/<li> in the generated lewislulu slide.
    """
    paras = []
    for p in text_frame.paragraphs:
        # Join runs so we keep the full paragraph text (runs split on formatting).
        text = "".join(run.text for run in p.runs)
        if not text:
            text = p.text  # fallback for paragraphs with no explicit runs
        text = text.strip()
        if not text:
            continue
        level = p.level if p.level is not None else 0
        paras.append({"text": text, "level": level})
    return paras


def extract_pptx(file_path, output_dir):
    prs = Presentation(file_path)

    images_dir = os.path.join(output_dir, "images")
    os.makedirs(images_dir, exist_ok=True)

    # Slide size in px — lets the agent reason about original aspect ratio.
    deck_meta = {
        "source": os.path.basename(file_path),
        "slide_width_px": emu_to_px(prs.slide_width),
        "slide_height_px": emu_to_px(prs.slide_height),
        "slide_count": len(prs.slides),
    }

    slides_data = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_data = {
            "number": slide_num,
            "title": "",
            "body": [],        # list of {text, level} paragraphs (bullets)
            "tables": [],      # list of {rows: [[cell, ...], ...]}
            "images": [],      # list of {path, width_px, height_px}
            "notes": "",
            "warnings": [],    # per-slide fidelity notes (charts, smartart, ...)
        }

        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None

        for shape in sorted(slide.shapes, key=shape_top):
            stype = shape.shape_type

            # --- Title ---
            if title_shape is not None and shape == title_shape:
                if shape.has_text_frame:
                    slide_data["title"] = shape.text_frame.text.strip()
                continue

            # --- Body text (any non-title text box) ---
            if shape.has_text_frame:
                paras = extract_paragraphs(shape.text_frame)
                slide_data["body"].extend(paras)
                continue

            # --- Tables ---
            if shape.has_table:
                rows = []
                for row in shape.table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                slide_data["tables"].append({"rows": rows})
                continue

            # --- Pictures ---
            if stype == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    ext = image.ext or "png"
                    idx = len(slide_data["images"]) + 1
                    name = f"slide{slide_num:02d}_img{idx}.{ext}"
                    path = os.path.join(images_dir, name)
                    with open(path, "wb") as fh:
                        fh.write(image.blob)
                    slide_data["images"].append({
                        "path": os.path.join("images", name).replace("\\", "/"),
                        "width_px": emu_to_px(shape.width),
                        "height_px": emu_to_px(shape.height),
                    })
                except Exception as e:
                    slide_data["warnings"].append(f"picture extract failed: {e}")
                continue

            # --- Things we can't faithfully carry to flat HTML: flag them ---
            if stype == MSO_SHAPE_TYPE.CHART:
                slide_data["warnings"].append(
                    "chart present — not extracted; rebuild as a Chart.js block "
                    "or static image"
                )
            elif stype == MSO_SHAPE_TYPE.GROUP:
                # Walk one level into groups for any text we can salvage.
                for sub in shape.shapes:
                    if sub.has_text_frame:
                        slide_data["body"].extend(extract_paragraphs(sub.text_frame))
                slide_data["warnings"].append(
                    "grouped shapes flattened — original layout/positioning lost"
                )
            elif stype == MSO_SHAPE_TYPE.MEDIA:
                slide_data["warnings"].append(
                    "embedded media (video/audio) — not extracted"
                )

        # --- Speaker notes ---
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            slide_data["notes"] = (notes_frame.text or "").strip()

        slides_data.append(slide_data)

    return deck_meta, slides_data


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("usage: python extract-pptx.py <input.pptx> [output_dir]\n")
        sys.exit(1)

    input_file = sys.argv[1]
    if not os.path.isfile(input_file):
        sys.stderr.write(f"error: file not found: {input_file}\n")
        sys.exit(1)

    if len(sys.argv) > 2:
        output_dir = sys.argv[2]
    else:
        stem = os.path.splitext(os.path.basename(input_file))[0]
        parent = os.path.dirname(os.path.abspath(input_file))
        output_dir = os.path.join(parent, f"{stem}-extract")

    os.makedirs(output_dir, exist_ok=True)

    deck_meta, slides = extract_pptx(input_file, output_dir)

    out = {"deck": deck_meta, "slides": slides}
    json_path = os.path.join(output_dir, "slides.json")
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, indent=2, ensure_ascii=False)

    print(f"extracted {deck_meta['slide_count']} slide(s) -> {json_path}")
    print(f"images    -> {os.path.join(output_dir, 'images')}/")
    for s in slides:
        bits = []
        bits.append(f"{len(s['body'])} para")
        if s["images"]:
            bits.append(f"{len(s['images'])} img")
        if s["tables"]:
            bits.append(f"{len(s['tables'])} table")
        if s["warnings"]:
            bits.append(f"{len(s['warnings'])} warning")
        title = s["title"] or "(no title)"
        print(f"  slide {s['number']:>2}: {title}  [{', '.join(bits)}]")

    # Surface warnings so the operator knows what may need manual rebuilding.
    flagged = [(s["number"], w) for s in slides for w in s["warnings"]]
    if flagged:
        print("\nfidelity warnings (rebuild these by hand):")
        for num, w in flagged:
            print(f"  slide {num}: {w}")


if __name__ == "__main__":
    main()
